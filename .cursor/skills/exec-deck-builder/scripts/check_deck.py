#!/usr/bin/env python3
"""pptx 交付前质检：溢出、越界、对比度、字号、东亚字体、重叠、信息过载。

为什么必须跑：这几类缺陷在代码里都看不出来，但投到大屏上每一个都很显眼。
文字被裁掉和灰底灰字是汇报现场最常见的两种事故。

用法：
    python3 check_deck.py deck.pptx            # 全量检查
    python3 check_deck.py deck.pptx --strict   # 有任何问题就返回非零退出码

退出码：0 无 error 级问题；1 有 error 级问题（--strict 下 warn 也算）。
"""

import os
import re
import sys

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import BASE_W, text_height, text_width, wrapped_lines  # noqa: E402

# 分级：error 必须修，warn 需要看一眼再决定
LEVELS = {
    "text_overflow": "error",
    "out_of_bounds": "error",
    "low_contrast": "error",
    "tiny_text": "error",
    "missing_ea_font": "error",
    "text_overlap": "warn",
    "tight_margin": "warn",
    "too_dense": "warn",
    "placeholder_left": "error",
}


def _lum(hexstr):
    vals = []
    for i in (0, 2, 4):
        c = int(hexstr[i:i + 2], 16) / 255.0
        vals.append(c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)
    return 0.2126 * vals[0] + 0.7152 * vals[1] + 0.0722 * vals[2]


def contrast(fg, bg):
    """WCAG 对比度，1:1 到 21:1。正文要求 >= 4.5，大号字 >= 3.0。"""
    a, b = _lum(fg), _lum(bg)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


def _solid_fill_rgb(obj):
    """取形状/背景的纯色填充 hex，取不到返回 None（渐变、图片、继承都算取不到）。"""
    try:
        fill = obj.fill
        if fill.type != 1:
            return None
        c = fill.fore_color
        if c.type != 1:
            return None
        return str(c.rgb).upper()
    except Exception:
        return None


def audit(path, verbose=True):
    prs = Presentation(path)
    pw = prs.slide_width / 914400
    ph = prs.slide_height / 914400
    scale = pw / BASE_W
    issues = []

    def add(sn, kind, msg):
        issues.append({"slide": sn, "kind": kind, "level": LEVELS.get(kind, "warn"),
                       "msg": msg})

    for sn, slide in enumerate(prs.slides, 1):
        page_bg = _solid_fill_rgb(slide.background) or "FFFFFF"
        # 记录带纯色填充的矩形，用于推断文字实际压在什么颜色上
        panels = []
        text_boxes = []

        for sh in slide.shapes:
            if sh.left is None or sh.top is None:
                continue
            x, y = sh.left / 914400, sh.top / 914400
            w, h = (sh.width or 0) / 914400, (sh.height or 0) / 914400
            rgb = _solid_fill_rgb(sh)
            if rgb and w * h > 0.02:
                panels.append((x, y, w, h, rgb))

            if x < -0.02 or y < -0.02 or x + w > pw + 0.02 or y + h > ph + 0.02:
                add(sn, "out_of_bounds",
                    "%s 超出画布 (%.2f,%.2f) %.2fx%.2f" % (sh.name, x, y, w, h))
            elif (x < 0.2 * scale or y < 0.2 * scale
                  or x + w > pw - 0.2 * scale or y + h > ph - 0.2 * scale):
                if sh.has_text_frame and sh.text_frame.text.strip():
                    add(sn, "tight_margin",
                        "%s 距页边不足 0.2\": (%.2f,%.2f)" % (sh.name, x, y))

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            txt = tf.text
            if not txt.strip():
                continue

            if re.search(r"lorem|ipsum|TODO|待填|xxx|XXX|\[插入", txt):
                add(sn, "placeholder_left", "%s 残留占位文字: %r" % (sh.name, txt[:30]))

            sizes, fgs, missing_ea = [], [], False
            for p in tf.paragraphs:
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    if r.font.size:
                        sizes.append(r.font.size.pt)
                    try:
                        if r.font.color and r.font.color.type == 1:
                            fgs.append(str(r.font.color.rgb).upper())
                    except Exception:
                        pass
                    if re.search(r"[\u4e00-\u9fff]", r.text):
                        if r.font._element.find(qn("a:ea")) is None:
                            missing_ea = True
            if missing_ea:
                add(sn, "missing_ea_font",
                    "%s 中文未设东亚字体，会回落主题字体: %r" % (sh.name, txt[:24]))
            if not sizes:
                continue

            design_sz = max(sizes) / scale
            if min(sizes) / scale < 8.5:
                add(sn, "tiny_text", "%s 字号折算后仅 %.1fpt: %r"
                    % (sh.name, min(sizes) / scale, txt[:24]))

            # 溢出：按估算行数与框内净高比较
            ml = (tf.margin_left or 0) / 914400
            mr = (tf.margin_right or 0) / 914400
            mt = (tf.margin_top or 0) / 914400
            mb = (tf.margin_bottom or 0) / 914400
            inner_w = (w - ml - mr) / scale
            inner_h = (h - mt - mb) / scale
            vert = tf._txBody.bodyPr.get("vert")
            if vert in ("vert", "vert270"):
                inner_w, inner_h = inner_h, inner_w   # 竖排时宽高互换
            need = text_height(txt, max(inner_w, 0.05), design_sz, 1.28)
            if need > inner_h * 1.14 + 0.02:
                add(sn, "text_overflow",
                    "%s 文本预计溢出 需%.2f\" 实%.2f\": %r"
                    % (sh.name, need, inner_h, txt[:24]))

            # 对比度：找覆盖该文字框且面积最小的色块作为实际背景
            bg = page_bg
            best = None
            for px, py, pw_, ph_, prgb in panels:
                if (px - 0.02 <= x and py - 0.02 <= y
                        and px + pw_ + 0.02 >= x + w and py + ph_ + 0.02 >= y + h):
                    a = pw_ * ph_
                    if best is None or a < best[0]:
                        best = (a, prgb)
            if best:
                bg = best[1]
            for fg in set(fgs):
                ratio = contrast(fg, bg)
                need_ratio = 3.0 if design_sz >= 18 else 4.5
                if ratio < need_ratio:
                    add(sn, "low_contrast",
                        "%s #%s 压在 #%s 上对比度仅 %.1f:1（需 %.1f）: %r"
                        % (sh.name, fg, bg, ratio, need_ratio, txt[:20]))

            # 单行文本按对齐方式收缩到文字真实占位，否则页脚注释（左对齐）和
            # 页码（右对齐）这类共处一行的元素会被误判成重叠。
            # 竖排文本的 align 作用在竖直方向，所以要收缩高度而不是宽度。
            ex, ey, ew, eh = x, y, w, h
            if wrapped_lines(txt, max(inner_w, 0.05), design_sz) == 1:
                real = min(text_width(txt, design_sz) * scale,
                           h if vert else w)
                al = tf.paragraphs[0].alignment
                aname = str(al) if al is not None else "LEFT"
                if vert == "vert270":
                    if "RIGHT" in aname:
                        eh = real
                    elif "CENTER" in aname:
                        ey, eh = y + (h - real) / 2, real
                    else:
                        ey, eh = y + h - real, real
                elif vert:
                    eh = real if "RIGHT" not in aname else real
                    ey = y if "RIGHT" not in aname else y + h - real
                elif "RIGHT" in aname:
                    ex, ew = x + w - real, real
                elif "CENTER" in aname:
                    ex, ew = x + (w - real) / 2, real
                else:
                    ew = real
            text_boxes.append((ex, ey, ew, eh, sh.name, txt))

        # 文字块互相重叠（同一页两段文字压在一起，读者会看到糊字）
        for i in range(len(text_boxes)):
            for j in range(i + 1, len(text_boxes)):
                ax, ay, aw, ah, an, _ = text_boxes[i]
                bx, by, bw, bh, bn, _ = text_boxes[j]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.06 * scale and oy > 0.06 * scale:
                    ov = ox * oy
                    if ov > 0.25 * min(aw * ah, bw * bh):
                        add(sn, "text_overlap",
                            "%s 与 %s 文本区重叠 %.2f in²" % (an, bn, ov))

        # 密度用字数衡量而不是元素数：复合页型（四象限 + 时间轴）元素天然多，
        # 但只要字少就依然清爽；真正劝退听众的是满页文字。
        body_chars = sum(len(t.replace(" ", "")) for _, _, _, _, _, t in text_boxes)
        if body_chars > 420:
            add(sn, "too_dense", "单页约 %d 字，建议拆页或精简到 300 字以内"
                % body_chars)
        elif len(slide.shapes) > 80:
            add(sn, "too_dense", "单页 %d 个元素，检查是否可以合并简化"
                % len(slide.shapes))

    errors = [i for i in issues if i["level"] == "error"]
    warns = [i for i in issues if i["level"] == "warn"]
    if verbose:
        print("检查 %s：%d 页" % (os.path.basename(path), len(prs.slides)))
        if not issues:
            print("全部通过：无溢出、无越界、对比度达标、字体正确。")
        else:
            for lvl, group in (("ERROR", errors), ("WARN", warns)):
                for it in group:
                    print("  [%s] 第%d页 %s — %s"
                          % (lvl, it["slide"], it["kind"], it["msg"]))
            print("合计 %d 个 error，%d 个 warn" % (len(errors), len(warns)))
    return issues


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    strict = "--strict" in sys.argv
    found = audit(sys.argv[1])
    bad = [i for i in found if i["level"] == "error" or strict]
    sys.exit(1 if bad else 0)
