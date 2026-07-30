"""2026 年 GMP 管理重点工作推进（质量体系版）。

内容取自 质量.pptx，风格沿用 GMP合规管理/2026年GMP管理重点工作推进.pptx：
深蓝标题带 + 白字、浅底圆角卡片 + 细边轻阴影、蓝色强调、箭头表达层级。

版面逻辑：一~四 为四根支柱并排，下方箭头汇聚到第五项 —— 四点共同达成
「建立常态化飞检就绪机制」。第五项内容待补充，页面上留出空位。

文字只做排版层面的梳理（去掉源文里的 markdown 星号、中英文数字间补空格、
把「集中于：」后紧跟列表的重复冒号去掉），语义与原文一致。

模板安全区按版式背景图实测：页眉止于 y 0.749"，页脚色条自 y 5.409" 起。
画布 10 x 5.625，deckkit 设计单位 13.333 x 7.5，故实际字号 = 设计字号 x 0.75。

运行：python3 build_quality_2026.py [输出路径.pptx]
"""

from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(REPO, ".cursor", "skills", "exec-deck-builder",
                               "scripts"))

from deckkit import (LIGHT, Deck, Rect, add_shadow, apply_font,  # noqa: E402
                     check_overflow, text_width, wrapped_lines)
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Pt  # noqa: E402
from lxml import etree  # noqa: E402

TEMPLATE = os.path.join(REPO, "质量.pptx")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    HERE, "2026年GMP管理重点工作推进.pptx")

NAVY = "132556"
BLUE = "1F4E86"
ACCENT = "1D42E1"
INK = "252D30"
MUTED = "55677D"
LINE = "C7D3E6"
CARD = "F6F9FC"
BAND = "EDF2FA"
CHIP = "E3EBF7"

LS = 1.20
LS_EST = LS * 1.22
BUL = 0.11

T = LIGHT.variant(
    name="asymchem-quality2026",
    bg="FFFFFF", bg_alt="F7FAFD", surface=CARD, surface_alt=BAND,
    hairline=LINE, ink=INK, ink_muted=MUTED, ink_on_accent="FFFFFF",
    primary=BLUE, secondary="2E5E9E", accent=ACCENT,
    good="1A7F50", warn="A34E0F", bad="C0392B", neutral="5E7288",
    accent_text=ACCENT, primary_text=NAVY, secondary_text="1F4E86",
    margin=0.40, margin_top=1.07, margin_bottom=0.35,
    font_cn="微软雅黑", font_en="Arial",
)

# 模板里可能有多个同名版式，按源页所用版式的 part 名精确取回
from pptx import Presentation  # noqa: E402

_probe = Presentation(TEMPLATE)
LAYOUT_PART = _probe.slides[0].slide_layout.part.partname
del _probe

deck = Deck(theme=T, template=TEMPLATE)
_target = None
for _m in deck.prs.slide_masters:
    for _lo in _m.slide_layouts:
        if _lo.part.partname == LAYOUT_PART:
            _target = _lo
            break
    if _target is not None:
        break
if _target is None:
    raise SystemExit("找不到源页所用版式：%s" % LAYOUT_PART)
deck._blank = _target

s = deck.slide()
SCALE = deck.scale
PW, PH = 10.0, 5.625
SAFE_TOP, SAFE_BOTTOM = 0.80, 5.34
SAFE_L, SAFE_R = 0.30, 9.70


def A(v: float) -> float:
    return v / SCALE


def R(x, y, w, h) -> Rect:
    return Rect(A(x), A(y), A(w), A(h))


# --------------------------------------------------------------------------
# 文本工具
# --------------------------------------------------------------------------

def set_bullet(paragraph, char="●", indent=BUL):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        node = pPr.find(qn(tag))
        if node is not None:
            pPr.remove(node)
    emu = int(indent * 914400)
    pPr.set("marL", str(emu))
    pPr.set("indent", str(-emu))
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    bc = etree.SubElement(pPr, qn("a:buChar"))
    bc.set("char", char)


def norm(lines):
    return [{"text": x} if isinstance(x, str) else x for x in lines]


def plain(item) -> str:
    if isinstance(item, str):
        return item
    if item.get("runs"):
        return "".join(g.get("text", "") for g in item["runs"])
    return item.get("text", "")


def block_height(lines, box_w, size, bullet_indent=BUL) -> float:
    total = 0.0
    for item in norm(lines):
        w = box_w - (bullet_indent if item.get("bullet") else 0.0)
        sz = size * item.get("scale", 1.0)
        total += wrapped_lines(plain(item), max(A(w), 0.05), A(sz)) \
            * sz * LS_EST / 72.0
        total += item.get("space_before", 0) / 72.0
    return total


def text(x, y, w, h, content, size, bold=False, color="ink", align="left",
         anchor="top", ls=LS, font_cn=None, font_en=None):
    return s.text(R(x, y, w, h), content, size=A(size), bold=bold, color=color,
                  align=align, anchor=anchor, line_spacing=ls,
                  font_cn=font_cn, font_en=font_en)


def rich(x, y, w, h, lines, size, lead=0.0, anchor="top", bullet_indent=BUL):
    box = s.raw.shapes.add_textbox(s._i(A(x)), s._i(A(y)),
                                   s._i(A(w)), s._i(A(h)))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP,
                          "middle": MSO_ANCHOR.MIDDLE}[anchor]
    first = True
    for item in norm(lines):
        is_first = first
        p = tf.paragraphs[0] if is_first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = LS
        sb = item.get("space_before", 0) + (0 if is_first else lead)
        if sb:
            p.space_before = Pt(sb)
        if item.get("bullet"):
            set_bullet(p, indent=bullet_indent)
        runs = item.get("runs") or [{"text": item.get("text", "")}]
        for g in runs:
            run = p.add_run()
            run.text = g.get("text", "")
            # apply_font 写的是最终 pt；画布即实际尺寸，所以直接给实际 pt
            apply_font(run.font, T, size=size * item.get("scale", 1.0),
                       bold=g.get("bold", item.get("bold", False)),
                       color=g.get("color", item.get("color", "ink")))
    return box


def lead_for(lines, box_w, size, box_h, cap=4.0):
    slack = box_h - block_height(lines, box_w, size)
    return max(0.0, min(cap, slack * 72.0 * 0.6 / max(len(lines) - 1, 1)))


# ==========================================================================
# 内容（取自 质量.pptx）
# ==========================================================================
PILLARS = [
    {
        "head": "一、工厂合规管理",
        "kind": "chips",
        "lead_chip": "数据完整性和计算机化系统（最高优先级）",
        "chips": ["生产", "QC", "库房", "设备", "培训", "常态化合规审计"],
    },
    {
        "head": "二、强化质量体系闭环运行，杜绝体系「纸面化」",
        "lines": [
            {"text": "监管现状：大量 483 缺陷集中于偏差/OOS/OOT 调查不深入、"
                     "CAPA 流于形式、变更控制失控、投诉处理不完善、"
                     "质量部门缺乏独立决策权。", "color": "ink_muted"},
            {"text": "核心落地举措：", "bold": True, "color": "accent",
             "space_before": 4},
            {"text": "1. 严格执行变更、偏差、OOS、CAPA、投诉、年度产品质量回顾"
                     "六大质量程序，坚持根本原因分析，不只用临时纠正代替长效预防；"},
            {"text": "2. QA 独立履职，拥有物料放行、产品放行否决权；"},
            {"text": "3. 定期内部模拟 FDA 体系审计，主动挖掘系统性隐患，"
                     "避免问题长期积累；"},
            {"text": "4. 所有质量活动证据完整留存，逻辑自洽、前后无矛盾。"},
        ],
    },
    {
        "head": "三、保障现场硬件状态、工艺与设施受控（现场直观核查重点）",
        "lines": [
            {"text": "核心落地举措：", "bold": True, "color": "accent"},
            {"text": "1. 维持设施、设备持续状态，验证文件与现场实际操作保持一致；"},
            {"text": "2. 设备预防性维护、仪器校准按期执行，状态标识清晰；"},
            {"text": "3. 洁净区人流物流分离、压差、温湿度、环境监测数据真实连续；"},
            {"text": "4. 物料分区隔离（待验/合格/不合格/留样），全链条可追溯；"
                     "清场管理、防止交叉污染。"},
        ],
    },
    {
        "head": "四、全员合规能力与质量文化建设，并打造 SME",
        "lines": [
            {"text": "核心落地举措：", "bold": True, "color": "accent"},
            {"text": "1. 持续开展 FDA-cGMP、数据完整性、飞检应答规范常态化培训，"
                     "不仅仅是理论，增加情景模拟；"},
            {"text": "2. 严格持证上岗，转岗、复工重新培训考核；"
                     "杜绝凭经验操作、擅自偏离 SOP；"},
            {"text": "3. 建立统一迎检规则：只回答提问、不主动延伸、不清楚不猜测、"
                     "主动查阅 SOP；"},
            {"text": "4. 营造主动上报异常的文化，杜绝隐瞒偏差、掩盖问题。"},
        ],
    },
]

GOAL_HEAD = "五、建立常态化飞检就绪机制"
GOAL_PLACEHOLDER = "（内容待补充）"


# ==========================================================================
# 页标题
# ==========================================================================
text(4.05, 0.18, 5.65, 0.34, "2026年GMP管理重点工作推进", 19.5, bold=True,
     color="primary_text", align="right", anchor="middle", font_cn="微软雅黑")
text(4.05, 0.53, 5.65, 0.22,
     "2026 GMP Governance Priorities & Implementation", 10.5, bold=True,
     color="secondary_text", align="right")


# ==========================================================================
# 版面：四根支柱并排 → 箭头汇聚 → 第五项目标条
# ==========================================================================
GOAL_H = 0.80
ARROW_H = 0.20
GAP = 0.14
PAD = 0.13
HEAD_H = 0.48

TOP = SAFE_TOP
PILLAR_H = SAFE_BOTTOM - GOAL_H - ARROW_H - 0.10 - TOP

# 栏宽按字数加权；"一"栏字少但要放标签块，给它一个下限
weights = []
for p in PILLARS:
    if p.get("kind") == "chips":
        n = len(p["lead_chip"]) + sum(len(c) for c in p["chips"]) + 40
    else:
        n = sum(len(plain(i)) for i in p["lines"])
    weights.append(n ** 0.6)
usable = SAFE_R - SAFE_L - GAP * (len(PILLARS) - 1)
widths = [usable * w / sum(weights) for w in weights]
FLOOR1 = 1.95                      # "一"栏要放得下最长那个标签块，且不挤成 3 行
if widths[0] < FLOOR1:
    extra = FLOOR1 - widths[0]
    rest = sum(widths[1:])
    widths = [FLOOR1] + [w - extra * w / rest for w in widths[1:]]

# 栏头统一字号：最长的那句也要排得进两行
head_size = 10.5
while head_size > 8.0:
    if all(wrapped_lines(p["head"], A(w - 2 * PAD), A(head_size)) <= 2
           for p, w in zip(PILLARS, widths)):
        break
    head_size -= 0.25

body_h = PILLAR_H - HEAD_H - 0.16
size = 9.6
while size > 7.0:
    ok = True
    for p, w in zip(PILLARS, widths):
        if p.get("kind") == "chips":
            continue
        if block_height(p["lines"], w - 2 * PAD, size) > body_h:
            ok = False
            break
    if ok:
        break
    size = round(size - 0.1, 1)


def draw_chips(col_x, col_w, y, h, data, csize):
    """一栏内的标签块：首个是最高优先级，做成高亮块；其余单列长条排布。

    单列而不是两列 —— 两列时每格只有 0.66" 宽，「常态化合规审计」会折行。
    """
    lead_h = 0.54
    box = R(col_x, y, col_w, lead_h)
    s.rect(box, fill=ACCENT, radius=0.05)
    s.text(box.inset(A(0.08), A(0.04)), data["lead_chip"],
           size=A(csize), bold=True, color="ink_on_accent", align="center",
           anchor="middle", line_spacing=1.14, font_cn="微软雅黑")

    cy = y + lead_h + 0.12
    n = len(data["chips"])
    gap = 0.06
    ch = min(0.34, (h - lead_h - 0.12 - gap * (n - 1)) / n)
    for i, name in enumerate(data["chips"]):
        chip = R(col_x, cy + i * (ch + gap), col_w, ch)
        s.rect(chip, fill=CHIP, line=LINE, line_w=0.8, radius=0.10)
        s.text(chip.inset(A(0.06), A(0.01)), name, size=A(csize),
               bold=True, color="primary_text", align="center",
               anchor="middle", font_cn="微软雅黑")


x = SAFE_L
centers = []
for p, w in zip(PILLARS, widths):
    col = R(x, TOP, w, PILLAR_H)
    card = s.rect(col, fill=CARD, line=LINE, line_w=1.0, radius=0.04)
    add_shadow(card, blur=7, dist=3, alpha=16)
    # 深蓝标题带 + 白字（沿用原页风格）
    s.rect(R(x, TOP, w, HEAD_H - 0.10), fill=NAVY, radius=0.05)
    s.rect(R(x, TOP + HEAD_H - 0.20, w, 0.10), fill=NAVY, radius=0)
    text(x + PAD, TOP, w - 2 * PAD, HEAD_H - 0.10, p["head"], head_size,
         bold=True, color="ink_on_accent", anchor="middle", ls=1.12,
         font_cn="微软雅黑")

    by = TOP + HEAD_H
    if p.get("kind") == "chips":
        draw_chips(x + PAD, w - 2 * PAD, by + 0.04, body_h, p, size + 0.4)
    else:
        rich(x + PAD, by, w - 2 * PAD, body_h, p["lines"], size,
             lead=lead_for(p["lines"], w - 2 * PAD, size, body_h, cap=3.0))
    centers.append(x + w / 2)
    x += w + GAP

# 四支柱 -> 目标：每栏正下方一个向下箭头，表达"共同达成"
ay = TOP + PILLAR_H + 0.02
for cx in centers:
    s.rect(R(cx - 0.11, ay, 0.22, ARROW_H), fill=ACCENT,
           shape=MSO_SHAPE.DOWN_ARROW, radius=0)

# 第五项：目标条。左侧深蓝标题块，右侧留一块白底待填区
GOAL_Y = SAFE_BOTTOM - GOAL_H
TAG_W5 = 2.62
goal = R(SAFE_L, GOAL_Y, SAFE_R - SAFE_L, GOAL_H)
s.rect(goal, fill=BAND, line=BLUE, line_w=1.4, radius=0.05)
s.rect(R(SAFE_L, GOAL_Y, TAG_W5, GOAL_H), fill=NAVY, radius=0.05)
s.rect(R(SAFE_L + TAG_W5 - 0.10, GOAL_Y, 0.10, GOAL_H), fill=NAVY, radius=0)
text(SAFE_L + 0.16, GOAL_Y, TAG_W5 - 0.32, GOAL_H, GOAL_HEAD, 12.0,
     bold=True, color="ink_on_accent", anchor="middle", ls=1.14,
     font_cn="微软雅黑")

slot = R(SAFE_L + TAG_W5 + 0.14, GOAL_Y + 0.09,
         SAFE_R - SAFE_L - TAG_W5 - 0.28, GOAL_H - 0.18)
s.rect(slot, fill="FFFFFF", line=LINE, line_w=1.0, radius=0.04)
s.text(slot, GOAL_PLACEHOLDER, size=A(9.5), color="ink_muted",
       align="center", anchor="middle", font_cn="微软雅黑")

deck.save(OUT)
print("已生成：%s（%d 页，%.2f x %.2f 英寸）"
      % (OUT, len(deck.slides), PW, PH))
print("栏宽：%s" % ["%.2f" % w for w in widths])
print("栏头 %.2fpt，正文 %.1fpt" % (head_size, size))
for p, w in zip(PILLARS, widths):
    if p.get("kind") == "chips":
        continue
    print("  %-22s 占高 %.2f\" / 可用 %.2f\""
          % (p["head"][:22], block_height(p["lines"], w - 2 * PAD, size),
             body_h))
check_overflow(OUT)
